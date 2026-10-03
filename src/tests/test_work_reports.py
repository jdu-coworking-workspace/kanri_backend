import calendar
from datetime import date, time
from io import BytesIO

import pytest
from openpyxl import load_workbook

from src.models.student import Student


def _auth(client, cookie):
    client.cookies.clear()
    client.cookies.set("access_token", cookie)


def _month_days(year, month, complete=True):
    _, last_day = calendar.monthrange(year, month)
    days = []
    for day in range(1, last_day + 1):
        work_date = date(year, month, day).isoformat()
        if not complete and day > 1:
            days.append({"work_date": work_date, "is_day_off": False})
            continue
        if day == 1:
            days.append(
                {
                    "work_date": work_date,
                    "is_day_off": True,
                    "description": "休み",
                }
            )
        elif day == 2:
            days.append(
                {
                    "work_date": work_date,
                    "is_day_off": False,
                    "start_time": "13:00",
                    "finish_time": "16:00",
                    "description": "night",
                }
            )
        else:
            days.append(
                {
                    "work_date": work_date,
                    "is_day_off": False,
                    "start_time": "09:00",
                    "finish_time": "12:00",
                    "description": "work",
                }
            )
    return days


def _payload(year, month, complete=True):
    return {
        "year": year,
        "month": month,
        "week_1": "week one",
        "days": _month_days(year, month, complete=complete),
    }


@pytest.fixture
def linked_student(test_db, student_user, sample_student) -> Student:
    sample_student.user_id = student_user.id
    test_db.commit()
    test_db.refresh(sample_student)
    return sample_student


class TestWorkReports:
    def test_student_submits_and_admin_reviews(
        self, client, linked_student, student_cookie, staff_cookie, admin_cookie, capsys
    ):
        today = date.today()
        year, month = today.year, today.month
        day_count = calendar.monthrange(year, month)[1]
        _auth(client, student_cookie)

        empty = client.get("/api/v1/reports/me")
        assert empty.status_code == 200, empty.text
        assert empty.json()["data"]["items"] == []
        assert empty.json()["data"]["can_add_current"] is True
        assert empty.json()["data"]["current_year"] == year
        assert empty.json()["data"]["current_month"] == month

        created = client.post("/api/v1/reports/me")
        assert created.status_code == 200, created.text
        body = created.json()["data"]
        report_id = body["id"]
        assert body["status"] == "draft"
        assert body["editable"] is True
        assert body["year"] == year
        assert body["month"] == month
        assert body["day_count"] == day_count
        assert body["student_code"] == linked_student.student_code
        assert client.post("/api/v1/reports/me").status_code == 409

        _auth(client, admin_cookie)
        hidden = client.get(f"/api/v1/reports/{report_id}")
        assert hidden.status_code == 403
        assert hidden.json()["detail"]["code"] == "REPORT_DRAFT_HIDDEN"
        _auth(client, student_cookie)

        incomplete = client.put(
            f"/api/v1/reports/me/{report_id}", json=_payload(year, month, complete=False)
        )
        assert incomplete.status_code == 200, incomplete.text
        denied = client.post(f"/api/v1/reports/me/{report_id}/submit")
        assert denied.status_code == 400
        assert denied.json()["detail"]["code"] == "REPORT_INCOMPLETE"

        zero_hours = _payload(year, month, complete=True)
        for day in zero_hours["days"]:
            if day["work_date"].endswith("-02"):
                day["start_time"] = "09:00"
                day["finish_time"] = "09:00"
        assert client.put(f"/api/v1/reports/me/{report_id}", json=zero_hours).status_code == 200
        zero_submit = client.post(f"/api/v1/reports/me/{report_id}/submit")
        assert zero_submit.status_code == 400
        assert zero_submit.json()["detail"]["code"] == "REPORT_INCOMPLETE"

        saved = client.put(
            f"/api/v1/reports/me/{report_id}", json=_payload(year, month, complete=True)
        )
        assert saved.status_code == 200, saved.text

        refreshed = client.get(f"/api/v1/reports/me/{report_id}")
        assert refreshed.status_code == 200, refreshed.text
        saved_body = refreshed.json()["data"]
        rest_day = next(day for day in saved_body["days"] if day["work_date"].endswith("-01"))
        assert rest_day["is_day_off"] is True
        assert rest_day["start_time"] is None
        assert rest_day["finish_time"] is None
        assert rest_day["description"] is None
        assert rest_day["duration_minutes"] == 0
        overnight = next(day for day in saved_body["days"] if day["work_date"].endswith("-02"))
        assert overnight["start_time"] == "13:00"
        assert overnight["finish_time"] == "16:00"
        assert overnight["description"] == "night"
        assert overnight["overnight"] is False
        assert overnight["duration_minutes"] == 180

        invalid = client.put(
            f"/api/v1/reports/me/{report_id}",
            json={
                "year": year,
                "month": month,
                "days": [{
                    "work_date": f"{year}-{month:02d}-02",
                    "start_time": "22:00",
                    "finish_time": "01:00",
                    "description": "night",
                    "is_day_off": False,
                }],
            },
        )
        assert invalid.status_code == 400
        assert invalid.json()["detail"]["code"] == "REPORT_TIME_ORDER"
        assert saved_body["total_minutes"] == (day_count - 1) * 180

        submitted = client.post(f"/api/v1/reports/me/{report_id}/submit")
        assert submitted.status_code == 200, submitted.text
        assert submitted.json()["data"]["status"] == "submitted"
        assert submitted.json()["data"]["editable"] is False

        locked = client.put(
            f"/api/v1/reports/me/{report_id}", json=_payload(year, month, complete=True)
        )
        assert locked.status_code == 400
        assert locked.json()["detail"]["code"] == "REPORT_LOCKED"

        listed = client.get("/api/v1/reports", params={"year": year, "month": month})
        assert listed.status_code == 403

        _auth(client, staff_cookie)
        staff_list = client.get("/api/v1/reports", params={"year": year, "month": month})
        assert staff_list.status_code == 200
        assert {"year": year, "month": month} in staff_list.json()["data"]["periods"]
        item = next(
            row for row in staff_list.json()["data"]["items"] if row["student_id"] == str(linked_student.id)
        )
        assert item["report_id"] == report_id
        assert item["status"] == "submitted"
        staff_detail = client.get(f"/api/v1/reports/{report_id}")
        assert staff_detail.status_code == 200
        assert client.post(f"/api/v1/reports/{report_id}/reject", json={"review_note": "no"}).status_code == 403
        assert client.get(f"/api/v1/reports/{report_id}/export").status_code == 403

        _auth(client, admin_cookie)
        rejected = client.post(
            f"/api/v1/reports/{report_id}/reject",
            json={"review_note": "時間を確認してください"},
        )
        assert rejected.status_code == 200, rejected.text
        assert rejected.json()["data"]["status"] == "rejected"
        assert rejected.json()["data"]["editable"] is True

        _auth(client, student_cookie)
        own_list = client.get("/api/v1/reports/me")
        assert own_list.status_code == 200
        own_item = own_list.json()["data"]["items"][0]
        assert own_item["status"] == "rejected"
        assert own_item["editable"] is True
        assert own_item["review_note"] == "時間を確認してください"
        assert own_list.json()["data"]["can_add_current"] is False

        resubmit = client.post(f"/api/v1/reports/me/{report_id}/submit")
        assert resubmit.status_code == 200, resubmit.text
        assert resubmit.json()["data"]["status"] == "submitted"
        assert resubmit.json()["data"]["review_note"] is None

        _auth(client, admin_cookie)
        accepted = client.post(
            f"/api/v1/reports/{report_id}/accept",
            json={"review_note": "確認しました", "stipend_amount": 24000},
        )
        assert accepted.status_code == 200, accepted.text
        assert accepted.json()["data"]["status"] == "approved"
        assert accepted.json()["data"]["stipend_amount"] == 24000
        assert client.post(
            f"/api/v1/reports/{report_id}/reject",
            json={"review_note": "late"},
        ).status_code == 400

        exported = client.get(f"/api/v1/reports/{report_id}/export")
        assert exported.status_code == 200
        assert "spreadsheetml" in exported.headers["content-type"]
        assert f"UZ240001_{year}-{month:02d}.xlsx" in exported.headers["content-disposition"]
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook.active
        assert sheet["C3"].value == "UZ240001"
        assert sheet["C4"].value == "Karimova Nilufar"
        assert sheet["B2"].value == "成果報告書（固定制）　勤務記録表"
        assert sheet["E7"].value == "=D7-C7"
        assert sheet["E38"].value == "=SUM(E7:E37)"
        assert sheet["C7"].value.hour == 0 and sheet["C7"].value.minute == 0
        assert sheet["F7"].value == "休み"
        assert sheet["C8"].value.hour == 13
        assert sheet["D8"].value.hour == 16
        assert sheet["C41"].value == "week one"
        assert sheet["I41"].value == 24000
        assert sheet["I47"].value == "承認"

        mailed = capsys.readouterr().out
        assert mailed.count("email jonatilindi") == 2
        assert "nilufar@test.com" in mailed
        assert f"{year}-{month:02d}" in mailed
        assert "rejected" in mailed
        assert "accepted" in mailed
        assert "時間を確認してください" in mailed
        assert "確認しました" in mailed

    def test_future_month_and_other_roles_cannot_edit_own_report(
        self, client, linked_student, student_cookie, admin_cookie
    ):
        _auth(client, student_cookie)
        missing_report = client.get("/api/v1/reports/me/00000000-0000-0000-0000-000000000000")
        assert missing_report.status_code == 404
        assert client.get("/api/v1/reports/me").json()["data"]["items"] == []

        _auth(client, admin_cookie)
        assert client.get("/api/v1/reports/me").status_code == 403
        missing = client.get("/api/v1/reports", params={"year": 2026, "month": 2})
        assert missing.status_code == 200
        assert {"year": 2026, "month": 2} not in missing.json()["data"]["periods"]
        assert {"year": date.today().year, "month": date.today().month} in missing.json()["data"]["periods"]
        row = next(
            item for item in missing.json()["data"]["items"] if item["student_id"] == str(linked_student.id)
        )
        assert row["report_id"] is None
        assert row["status"] is None


class TestSeikaHoukokushoExcel:
    def test_fills_official_template_layout(self):
        from src.utils.report_excel import build_monthly_report_xlsx

        year, month = 2026, 2
        last_day = calendar.monthrange(year, month)[1]
        days = []
        for day in range(1, last_day + 1):
            work_date = date(year, month, day).isoformat()
            if day == 1:
                days.append({"work_date": work_date, "is_day_off": True})
            else:
                days.append(
                    {
                        "work_date": work_date,
                        "is_day_off": False,
                        "start_time": "09:00",
                        "finish_time": "12:00",
                        "description": "work",
                    }
                )
        content = build_monthly_report_xlsx(
            {
                "year": year,
                "month": month,
                "student_code": "2323123",
                "student_name": "武座",
                "status": "submitted",
                "stipend_amount": None,
                "week_2": "second week",
                "days": days,
            }
        )
        sheet = load_workbook(BytesIO(content)).active
        assert sheet["B2"].value == "成果報告書（固定制）　勤務記録表"
        assert sheet["C3"].value == "2323123"
        assert sheet["C4"].value == "武座"
        assert sheet["B6"].value == "日付"
        assert sheet["B7"].value.date() == date(year, month, 1)
        assert sheet["B34"].value.date() == date(year, month, 28)
        assert sheet["B35"].value is None
        assert sheet["F7"].value == "休み"
        assert sheet["C8"].value.hour == 9
        assert sheet["F8"].value == "work"
        assert sheet["C44"].value == "second week"
        assert sheet["I47"].value is None
        assert sheet["E7"].value == "=D7-C7"
        assert sheet["E38"].value == "=SUM(E7:E37)"



class TestSeikaHoukokushoBulkExport:
    def test_period_archive_and_draft_download(
        self, client, admin_cookie, staff_cookie, test_db, linked_student
    ):
        from io import BytesIO as _BytesIO
        from zipfile import ZipFile

        from src.models.work_report import DailyReport, MonthlyReport, ReportStatus

        year, month = date.today().year, date.today().month
        _auth(client, admin_cookie)
        last_day = calendar.monthrange(year, month)[1]
        draft = MonthlyReport(
            student_id=linked_student.id,
            year=year,
            month=month,
            status=ReportStatus.DRAFT,
            week_2="second week",
        )
        test_db.add(draft)
        test_db.flush()
        for day_number in range(1, last_day + 1):
            test_db.add(
                DailyReport(
                    monthly_report_id=draft.id,
                    work_date=date(year, month, day_number),
                    is_day_off=day_number == 1,
                    start_time=None if day_number == 1 else time(9, 0),
                    finish_time=None if day_number == 1 else time(17, 0),
                    description=None if day_number == 1 else "work",
                )
            )
        test_db.commit()

        items = client.get(f"/api/v1/reports?year={year}&month={month}").json()["data"]["items"]
        downloadable = [item for item in items if item["report_id"]]
        assert downloadable, "draft row must be downloadable by an admin"
        assert any(item["status"] == "draft" for item in downloadable)

        _auth(client, staff_cookie)
        assert client.get("/api/v1/reports/export", params={"year": year, "month": month}).status_code == 403
        assert client.get(f"/api/v1/reports/{downloadable[0]['report_id']}/export").status_code == 403

        _auth(client, admin_cookie)
        archive = client.get("/api/v1/reports/export", params={"year": year, "month": month})
        assert archive.status_code == 200, archive.text
        assert archive.headers["content-type"] == "application/zip"
        assert f"reports_{year}-{month:02d}.zip" in archive.headers["content-disposition"]

        with ZipFile(_BytesIO(archive.content)) as bundle:
            names = bundle.namelist()
            assert len(names) == len(downloadable)
            for name in names:
                assert name.endswith(".xlsx")
                sheet = load_workbook(_BytesIO(bundle.read(name))).active
                assert sheet["B2"].value == "成果報告書（固定制）　勤務記録表"

        empty = client.get("/api/v1/reports/export", params={"year": 2019, "month": 1})
        assert empty.status_code == 404
        assert empty.json()["detail"]["code"] == "REPORT_EXPORT_EMPTY"
