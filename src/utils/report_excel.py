import calendar
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path

# pyrefly: ignore [missing-import]
from openpyxl import load_workbook

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "assets" / "seika_houkokusho.xlsx"
DAY_START_ROW = 7
MAX_DAY_SLOTS = 31
WEEK_CELLS = {
    1: "C41",
    2: "C44",
    3: "C47",
    4: "C50",
    5: "C53",
}


def _as_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _as_time(value) -> time:
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if not value:
        return time(0, 0)
    parts = str(value).strip().split(":")
    return time(int(parts[0]), int(parts[1]) if len(parts) > 1 else 0)


def set_cell(sheet, row: int, column: int, value) -> None:
    """Worksheet.cell() silently ignores a None value, so assign explicitly."""
    sheet.cell(row=row, column=column).value = value


def build_monthly_report_xlsx(report: dict) -> bytes:
    """Fill the official 成果報告書 template the team already uses."""
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Excel template missing: {TEMPLATE_PATH}")

    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook.active
    year = int(report["year"])
    month = int(report["month"])
    last_day = calendar.monthrange(year, month)[1]
    days_by_date = {_as_date(day["work_date"]): day for day in report.get("days") or []}

    sheet["C3"] = report.get("student_code") or ""
    sheet["F3"] = report.get("student_code") or ""
    sheet["C4"] = report.get("student_name") or ""

    for offset in range(MAX_DAY_SLOTS):
        row = DAY_START_ROW + offset
        day_number = offset + 1
        if day_number > last_day:
            set_cell(sheet, row, 2, None)
            set_cell(sheet, row, 3, time(0, 0))
            set_cell(sheet, row, 4, time(0, 0))
            set_cell(sheet, row, 6, None)
            continue

        work_date = date(year, month, day_number)
        day = days_by_date.get(work_date) or {}
        set_cell(sheet, row, 2, work_date)
        if day.get("is_day_off"):
            set_cell(sheet, row, 3, time(0, 0))
            set_cell(sheet, row, 4, time(0, 0))
            set_cell(sheet, row, 6, "休み")
        else:
            set_cell(sheet, row, 3, _as_time(day.get("start_time")))
            set_cell(sheet, row, 4, _as_time(day.get("finish_time")))
            set_cell(sheet, row, 6, day.get("description") or None)

    for week, cell in WEEK_CELLS.items():
        sheet[cell] = report.get(f"week_{week}") or None

    sheet["I41"] = report.get("stipend_amount")
    if report.get("status") == "approved":
        sheet["I47"] = "承認"
    elif report.get("status") == "rejected":
        sheet["I47"] = "承認不可"
    else:
        sheet["I47"] = None

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
